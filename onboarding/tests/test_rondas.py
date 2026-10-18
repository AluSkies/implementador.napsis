"""Validar la documentación o devolvérsela al cliente: estados nuevos y rondas de revisión."""

from django.contrib.auth.models import User
from django.core import mail
from django.urls import reverse

from automatizacion.actores import Actor
from automatizacion.models import EventoImplementacion
from clients.models import Implementacion, RondaRevision
from clients.services import filtrar_y_ordenar_implementaciones
from onboarding import acciones
from onboarding.tests.test_wizard import BaseWizard, _url

CLIENTE = Actor.cliente("cliente_a")
ANA = Actor.implementador("ana")


class EnviadaBase(BaseWizard):
    """Una implementación completa y ya enviada por el cliente."""

    def setUp(self):
        super().setUp()
        self.implementacion.mail_contacto = "cliente@empresa.com"
        self.implementacion.save()
        self.completar_todo()
        acciones.enviar_relevamiento(CLIENTE, self.implementacion, None)
        self.implementacion.refresh_from_db()
        mail.outbox.clear()

    def devolver(self, motivo="Faltan los legajos del personal nuevo.", observaciones=None):
        return acciones.devolver_al_cliente(ANA, self.implementacion, motivo, observaciones)


class EstadosVisualesTests(EnviadaBase):
    def test_recien_enviada_espera_validacion(self):
        self.assertEqual(self.implementacion.estado_visual, Implementacion.VISUAL_COMPLETADO_CLIENTE)
        self.assertEqual(self.implementacion.estado_visual_label, "Documentación completa - Pendiente de validación")

    def test_al_validarla_queda_pendiente_de_configurar(self):
        acciones.aprobar_para_configurar(ANA, self.implementacion)
        self.assertEqual(self.implementacion.estado_visual, Implementacion.VISUAL_VALIDADA)
        self.assertEqual(self.implementacion.estado_visual_label, "Validada - Pendiente de configurar")

    def test_el_filtro_distingue_pendiente_de_validacion_y_validada(self):
        acciones.aprobar_para_configurar(ANA, self.implementacion)
        validada = Implementacion.objects.filter(Implementacion.filtro_visual(Implementacion.VISUAL_VALIDADA))
        sin_validar = Implementacion.objects.filter(Implementacion.filtro_visual(Implementacion.VISUAL_COMPLETADO_CLIENTE))
        self.assertEqual(list(validada), [self.implementacion])
        self.assertEqual(list(sin_validar), [])

    def test_el_listado_ordena_los_estados_en_su_orden_natural(self):
        otra = Implementacion.objects.create(cliente=self.cliente)
        acciones.aprobar_para_configurar(ANA, self.implementacion)
        orden = filtrar_y_ordenar_implementaciones(Implementacion.objects.all(), {"ordenar": "estado"})
        self.assertEqual(list(orden), [otra, self.implementacion])


class AprobarTests(EnviadaBase):
    def test_deja_la_fecha_y_el_evento(self):
        resultado = acciones.aprobar_para_configurar(ANA, self.implementacion)
        self.assertTrue(resultado.ok)
        self.assertIsNotNone(self.implementacion.validada_en)
        evento = EventoImplementacion.objects.get(accion="estado.validada")
        self.assertEqual((evento.actor_tipo, evento.actor_ref), ("implementador", "ana"))

    def test_no_se_valida_lo_que_el_cliente_no_envio(self):
        self.implementacion.enviado = None
        self.implementacion.save()
        resultado = acciones.aprobar_para_configurar(ANA, self.implementacion)
        self.assertIsNone(self.implementacion.validada_en)
        self.assertEqual(len(resultado.avisos), 1)
        self.assertFalse(EventoImplementacion.objects.filter(accion="estado.validada").exists())

    def test_no_se_valida_dos_veces(self):
        acciones.aprobar_para_configurar(ANA, self.implementacion)
        acciones.aprobar_para_configurar(ANA, self.implementacion)
        self.assertEqual(EventoImplementacion.objects.filter(accion="estado.validada").count(), 1)

    def test_si_despues_del_envio_falta_algo_no_se_valida(self):
        self.implementacion.empleados.all().delete()
        resultado = acciones.aprobar_para_configurar(ANA, self.implementacion)
        self.assertIsNone(self.implementacion.validada_en)
        self.assertIn("No se puede validar", resultado.avisos[0].texto)

    def test_reabrir_borra_la_validacion(self):
        acciones.aprobar_para_configurar(ANA, self.implementacion)
        acciones.reabrir(ANA, self.implementacion)
        self.assertIsNone(self.implementacion.validada_en)
        self.assertEqual(self.implementacion.estado_visual, Implementacion.VISUAL_A_COMPLETAR)


class DevolverTests(EnviadaBase):
    def test_abre_la_ronda_y_la_saca_del_envio(self):
        resultado = self.devolver(observaciones={"nomina": "Falta el personal de la sucursal Rosario."})
        self.assertTrue(resultado.ok)
        self.implementacion.refresh_from_db()
        self.assertIsNone(self.implementacion.enviado)
        self.assertEqual(self.implementacion.estado_visual, Implementacion.VISUAL_A_COMPLETAR)
        self.assertEqual(self.implementacion.marca_de_devolucion, "Devuelta, ronda 1")
        ronda = self.implementacion.rondas.get()
        self.assertEqual((ronda.numero, ronda.devuelta_por), (1, "ana"))
        self.assertEqual(ronda.observaciones, {"nomina": "Falta el personal de la sucursal Rosario."})

    def test_tambien_se_puede_devolver_una_ya_validada(self):
        acciones.aprobar_para_configurar(ANA, self.implementacion)
        self.devolver()
        self.implementacion.refresh_from_db()
        self.assertIsNone(self.implementacion.validada_en)

    def test_sin_motivo_no_se_devuelve(self):
        resultado = self.devolver(motivo="   ")
        self.implementacion.refresh_from_db()
        self.assertIsNotNone(self.implementacion.enviado)
        self.assertEqual(self.implementacion.rondas.count(), 0)
        self.assertEqual(len(resultado.avisos), 1)

    def test_lo_que_no_se_envio_no_se_devuelve(self):
        self.devolver()
        self.devolver(motivo="otra vez")  # ya está en manos del cliente
        self.assertEqual(self.implementacion.rondas.count(), 1)

    def test_descarta_observaciones_vacias_o_de_pasos_que_no_existen(self):
        self.devolver(observaciones={"nomina": "  ", "roles": "Falta el liquidador.", "inventado": "x"})
        self.assertEqual(self.implementacion.rondas.get().observaciones, {"roles": "Falta el liquidador."})

    def test_deja_el_evento_con_el_motivo(self):
        self.devolver()
        evento = EventoImplementacion.objects.get(accion="ronda.devuelta")
        self.assertEqual(evento.datos["ronda"], 1)
        self.assertEqual(evento.datos["motivo"], "Faltan los legajos del personal nuevo.")
        self.assertTrue(evento.post_envio)

    def test_le_avisa_al_cliente_por_mail(self):
        self.devolver(observaciones={"roles": "Falta el liquidador."})
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].to, ["cliente@empresa.com"])
        self.assertIn("Faltan los legajos del personal nuevo.", mail.outbox[0].body)
        self.assertIn("Falta el liquidador.", mail.outbox[0].body)
        self.assertIn(self.implementacion.token, mail.outbox[0].body)

    def test_si_el_mail_falla_queda_devuelta_igual(self):
        from unittest import mock

        with mock.patch("onboarding.acciones.notificaciones.notificar_cliente_devuelta", side_effect=OSError), \
                self.assertLogs("onboarding.acciones", level="ERROR"):
            resultado = self.devolver()
        self.assertEqual(self.implementacion.rondas.count(), 1)
        self.assertEqual(resultado.avisos[-1].texto, "Quedó devuelta, pero no se pudo avisarle al cliente por mail.")
        self.assertTrue(EventoImplementacion.objects.filter(accion="mail.fallido").exists())


class RespuestaDelClienteTests(EnviadaBase):
    def test_al_volver_a_enviar_se_cierra_la_ronda(self):
        self.devolver()
        acciones.enviar_relevamiento(CLIENTE, self.implementacion, None)
        self.implementacion.refresh_from_db()
        ronda = self.implementacion.rondas.get()
        self.assertEqual(ronda.respondida_en, self.implementacion.enviado)
        self.assertEqual(self.implementacion.marca_de_devolucion, "")
        self.assertEqual(self.implementacion.estado_visual, Implementacion.VISUAL_COMPLETADO_CLIENTE)
        envio = EventoImplementacion.objects.filter(accion="envio.confirmado").last()
        self.assertEqual(envio.datos["ronda"], 1)

    def test_una_segunda_devolucion_es_la_ronda_2(self):
        self.devolver()
        acciones.enviar_relevamiento(CLIENTE, self.implementacion, None)
        self.devolver(motivo="Falta el modelo de recibo firmado.")
        self.assertEqual(self.implementacion.marca_de_devolucion, "Devuelta, ronda 2")
        self.assertEqual(list(self.implementacion.rondas.values_list("numero", flat=True)), [1, 2])

    def test_el_cliente_puede_editar_de_nuevo_y_ve_lo_que_le_pidieron(self):
        self.devolver(observaciones={"roles": "Falta el liquidador."})
        pagina = self.client.get(_url(self.implementacion, 5))  # roles
        self.assertEqual(pagina.status_code, 200)
        self.assertContains(pagina, "ronda 1")
        self.assertContains(pagina, "Faltan los legajos del personal nuevo.")
        self.assertContains(pagina, "Falta el liquidador.")

    def test_la_observacion_de_otro_paso_no_se_muestra(self):
        self.devolver(observaciones={"roles": "Falta el liquidador."})
        self.assertNotContains(self.client.get(_url(self.implementacion, 1)), "Falta el liquidador.")

    def test_sin_devolucion_no_hay_cartel(self):
        self.implementacion.enviado = None
        self.implementacion.save()
        self.assertNotContains(self.client.get(_url(self.implementacion, 1)), "te devolvió el relevamiento")

    def test_mis_implementaciones_muestra_la_marca(self):
        self.devolver()
        self.assertContains(self.client.get(reverse("onboarding:mis_implementaciones")), "Devuelta, ronda 1")


class PanelTests(EnviadaBase):
    def setUp(self):
        super().setUp()
        self.client.force_login(User.objects.create_user("ana", password="x", is_staff=True))
        self.url = reverse("dashboard:implementacion_detalle", kwargs={"pk": self.implementacion.pk})

    def test_la_ficha_ofrece_validar_y_devolver(self):
        pagina = self.client.get(self.url)
        self.assertContains(pagina, "Validar documentación y dejar lista para configurar")
        self.assertContains(pagina, "Devolver al cliente")
        self.assertContains(pagina, 'name="obs_nomina"')

    def test_validar_desde_la_ficha(self):
        self.assertEqual(self.client.post(self.url, {"accion": "aprobar_para_configurar"}).status_code, 302)
        self.implementacion.refresh_from_db()
        self.assertEqual(self.implementacion.estado_visual, Implementacion.VISUAL_VALIDADA)
        pagina = self.client.get(self.url)
        self.assertContains(pagina, "Validada - Pendiente de configurar")
        self.assertNotContains(pagina, "Validar documentación y dejar lista")

    def test_devolver_desde_la_ficha_con_observaciones(self):
        self.client.post(self.url, {
            "accion": "devolver_al_cliente", "motivo": "Revisá la nómina.", "obs_nomina": "Faltan 3 personas.", "obs_roles": "",
        })
        ronda = RondaRevision.objects.get()
        self.assertEqual((ronda.motivo, ronda.observaciones, ronda.devuelta_por), ("Revisá la nómina.", {"nomina": "Faltan 3 personas."}, "ana"))
        pagina = self.client.get(self.url)
        self.assertContains(pagina, "Devuelta, ronda 1")
        self.assertContains(pagina, "Esperando al cliente")

    def test_devolver_sin_motivo_avisa_y_no_hace_nada(self):
        respuesta = self.client.post(self.url, {"accion": "devolver_al_cliente", "motivo": ""}, follow=True)
        self.assertContains(respuesta, "Contale al cliente qué tiene que corregir")
        self.assertEqual(RondaRevision.objects.count(), 0)

    def test_el_listado_filtra_por_validada_y_muestra_la_marca(self):
        acciones.aprobar_para_configurar(ANA, self.implementacion)
        home = reverse("dashboard:home")
        self.assertContains(self.client.get(home, {"estado": "validada"}), "Empresa A</a></td>")
        self.assertContains(self.client.get(home, {"estado": "completado_cliente"}), "No hay implementaciones que coincidan")
        acciones.devolver_al_cliente(ANA, self.implementacion, "Ajustar")
        self.assertContains(self.client.get(home), "Devuelta, ronda 1")

    def test_el_reporte_incluye_validacion_y_rondas(self):
        import openpyxl

        from dashboard.reportes import generar_reporte_implementaciones_xlsx

        acciones.devolver_al_cliente(ANA, self.implementacion, "Ajustar")
        hoja = openpyxl.load_workbook(generar_reporte_implementaciones_xlsx()).active
        fila = {hoja.cell(1, c).value: hoja.cell(2, c).value for c in range(1, hoja.max_column + 1)}
        self.assertEqual(fila["Estado"], "Devuelta, ronda 1")
        self.assertEqual(fila["Rondas de devolución"], 1)
