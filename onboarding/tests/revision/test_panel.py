"""Validar y devolver desde la ficha, el listado y el reporte."""

from dashboard.tests.base import crear_supervisor
from django.contrib.auth.models import User
from django.urls import reverse

from automatizacion.actores import Actor
from clients.models import Implementacion, RondaRevision
from onboarding import acciones
from onboarding.tests.base import EnviadaBase

ANA = Actor.implementador("ana")


class PanelTests(EnviadaBase):
    def setUp(self):
        super().setUp()
        self.client.force_login(crear_supervisor("ana"))
        self.url = reverse("dashboard:implementacion_detalle", kwargs={"pk": self.implementacion.pk})

    def test_la_ficha_ofrece_validar_y_devolver(self):
        pagina = self.client.get(self.url)
        self.assertContains(pagina, "Validar documentación y dejar lista para configurar")
        self.assertContains(pagina, "Devolver al cliente")
        self.assertContains(pagina, 'name="obs_nomina"')

    def test_validar_desde_la_ficha(self):
        self.assertEqual(self.client.post(self.url, {"accion": "aprobar_para_configurar"}).status_code, 302)
        self.implementacion.refresh_from_db()
        self.assertEqual(self.implementacion.estado_visual, Implementacion.VISUAL_VALIDADA)
        pagina = self.client.get(self.url)
        self.assertContains(pagina, "Validada - Pendiente de configurar")
        self.assertNotContains(pagina, "Validar documentación y dejar lista")

    def test_devolver_desde_la_ficha_con_observaciones(self):
        self.client.post(self.url, {
            "accion": "devolver_al_cliente", "motivo": "Revisá la nómina.", "obs_nomina": "Faltan 3 personas.", "obs_roles": "",
        })
        ronda = RondaRevision.objects.get()
        self.assertEqual((ronda.motivo, ronda.observaciones, ronda.devuelta_por), ("Revisá la nómina.", {"nomina": "Faltan 3 personas."}, "ana"))
        pagina = self.client.get(self.url)
        self.assertContains(pagina, "Devuelta, ronda 1")
        self.assertContains(pagina, "Esperando al cliente")

    def test_devolver_sin_motivo_avisa_y_no_hace_nada(self):
        respuesta = self.client.post(self.url, {"accion": "devolver_al_cliente", "motivo": ""}, follow=True)
        self.assertContains(respuesta, "Contale al cliente qué tiene que corregir")
        self.assertEqual(RondaRevision.objects.count(), 0)

    def test_el_listado_filtra_por_validada_y_muestra_la_marca(self):
        acciones.aprobar_para_configurar(ANA, self.implementacion)
        home = reverse("dashboard:home")
        self.assertContains(self.client.get(home, {"estado": "validada"}), "Empresa A</a></td>")
        self.assertContains(self.client.get(home, {"estado": "completado_cliente"}), "No hay implementaciones que coincidan")
        acciones.devolver_al_cliente(ANA, self.implementacion, "Ajustar")
        self.assertContains(self.client.get(home), "Devuelta, ronda 1")

    def test_el_reporte_incluye_validacion_y_rondas(self):
        import openpyxl

        from dashboard.reportes import generar_reporte_implementaciones_xlsx

        acciones.devolver_al_cliente(ANA, self.implementacion, "Ajustar")
        hoja = openpyxl.load_workbook(generar_reporte_implementaciones_xlsx()).active
        fila = {hoja.cell(1, c).value: hoja.cell(2, c).value for c in range(1, hoja.max_column + 1)}
        self.assertEqual(fila["Estado"], "Devuelta, ronda 1")
        self.assertEqual(fila["Rondas de devolución"], 1)


class AvisoAlPerderLaValidacionTests(EnviadaBase):
    """Si el equipo cambia algo de una implementación validada, la ficha se lo avisa."""

    AVISO = "volvió a «Pendiente de validación»"

    def setUp(self):
        super().setUp()
        self.client.force_login(crear_supervisor("ana"))
        self.url = reverse("dashboard:implementacion_detalle", kwargs={"pk": self.implementacion.pk})
        acciones.aprobar_para_configurar(ANA, self.implementacion)

    def datos(self, **cambios):
        self.implementacion.refresh_from_db()
        base = {"accion": "actualizar_datos", "razon_social": "Empresa A S.A.", "cuits": "30-50001091-2",
                "telefono": self.implementacion.telefono, "mail_contacto": self.implementacion.mail_contacto}
        return self.client.post(self.url, {**base, **cambios}, follow=True)

    def test_cambiar_un_dato_avisa_que_hay_que_volver_a_validar(self):
        respuesta = self.datos(telefono="011-5555-6666")
        self.assertContains(respuesta, self.AVISO)
        self.assertContains(respuesta, "Documentación completa - Pendiente de validación")

    def test_guardar_lo_mismo_no_avisa_porque_no_se_perdio_nada(self):
        self.assertNotContains(self.datos(), self.AVISO)

    def test_un_comentario_no_avisa(self):
        respuesta = self.client.post(self.url, {"accion": "agregar_comentario", "texto": "nota"}, follow=True)
        self.assertNotContains(respuesta, self.AVISO)

    def test_devolver_al_cliente_no_repite_el_aviso(self):
        respuesta = self.client.post(self.url, {"accion": "devolver_al_cliente", "motivo": "Ajustar"}, follow=True)
        self.assertNotContains(respuesta, self.AVISO)
        self.assertContains(respuesta, "Devuelta al cliente (ronda 1)")

    def test_si_no_estaba_validada_tampoco_avisa(self):
        self.datos(telefono="011-1")  # pierde la validación
        self.assertNotContains(self.datos(telefono="011-2"), self.AVISO)
