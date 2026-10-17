"""Parser/validador de la plantilla masiva de empleados y generador de la plantilla vacía."""

import io
import numbers
import re
import unicodedata
from collections import namedtuple
from datetime import date, datetime, timedelta

import openpyxl
import pandas as pd
from django.core.exceptions import ValidationError
from django.core.validators import validate_email

from clients.xlsx import escribir_celda

# (header visible en la plantilla, campo del modelo Empleado, obligatorio, tipo)
COLUMNAS = [
    ("CUIL", "cuil", True, "texto"),
    ("DNI", "dni", False, "texto"),
    ("Nombre", "nombre", True, "texto"),
    ("Apellido", "apellido", True, "texto"),
    ("Legajo", "legajo", False, "texto"),
    ("Fecha Ingreso (DD/MM/AAAA)", "fecha_ingreso", False, "fecha"),
    ("Mail", "mail", False, "mail"),
    ("Sexo (F/M)", "sexo", False, "sexo"),
    ("Fecha Nacimiento (DD/MM/AAAA)", "fecha_nacimiento", False, "fecha"),
    ("Lugar de Nacimiento", "lugar_nacimiento", False, "texto"),
    ("Nacionalidad", "nacionalidad", False, "texto"),
    ("Estado Civil", "estado_civil", False, "texto"),
    ("Domicilio", "domicilio", False, "texto"),
    ("Planta/Vivienda", "planta_vivienda", False, "texto"),
    ("Codigo Postal", "codigo_postal", False, "texto"),
    ("Localidad", "localidad", False, "texto"),
    ("Telefono", "telefono", False, "texto"),
    ("Obra Social", "obra_social", False, "texto"),
    ("Centro de Costo", "centro_costo", False, "texto"),
    ("Puesto", "puesto", False, "texto"),
    ("Gerencia", "gerencia", False, "texto"),
    ("Sector", "sector", False, "texto"),
    ("Funcion", "funcion", False, "texto"),
    ("Es Supervisor (SI/NO)", "es_supervisor", False, "booleano"),
    ("Home Office (SI/NO)", "home_office", False, "booleano"),
    ("Reporta A (CUIL)", "reporta_a", False, "cuil_ref"),
    ("Aprobador (CUIL)", "aprobador", False, "cuil_ref"),
]

CUIL_RE = re.compile(r"^\d{11}$")

# Más filas que esto en un solo archivo: se pide dividirlo (la carga corre dentro del pedido web).
MAX_FILAS_IMPORTACION = 5000

ResultadoLectura = namedtuple("ResultadoLectura", "filas errores columnas_ignoradas")


class ArchivoInvalido(Exception):
    """El archivo no se puede procesar. El mensaje está pensado para mostrárselo a quien lo subió."""


def normalizar_texto(valor):
    """Mayúsculas, sin espacios de más, sin tildes — mismo criterio usado en la plantilla Mondezza."""
    if valor is None:
        return ""
    texto = str(valor).strip().upper()
    texto = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", texto)


def normalizar_cuil(valor):
    if valor is None:
        return ""
    return re.sub(r"\D", "", str(valor))


def _armar_hoja_con_headers():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Empleados"
    for col_idx, (header, *_rest) in enumerate(COLUMNAS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = openpyxl.styles.Font(bold=True)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = max(18, len(header) + 2)
    return wb, ws


def generar_plantilla_xlsx():
    """Genera un .xlsx vacío con las columnas esperadas, para que el cliente lo descargue y complete."""
    wb, _ws = _armar_hoja_con_headers()
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generar_nomina_xlsx(empleados):
    """Genera un .xlsx con la nómina ya cargada de una implementación, para que el
    equipo de Napsis la descargue (mismas columnas que la plantilla vacía)."""
    wb, ws = _armar_hoja_con_headers()
    for fila_idx, empleado in enumerate(empleados, start=2):
        for col_idx, (_header, campo, _obligatorio, tipo) in enumerate(COLUMNAS, start=1):
            if campo in ("reporta_a", "aprobador"):
                relacionado = getattr(empleado, campo)
                valor = relacionado.cuil if relacionado else ""
            else:
                valor = getattr(empleado, campo)
                if tipo == "booleano":
                    valor = "SI" if valor else "NO"
                elif tipo == "fecha" and valor:
                    valor = valor.strftime("%d/%m/%Y")
                elif valor is None:
                    valor = ""
            escribir_celda(ws, fila_idx, col_idx, valor)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# --- Lectura del archivo ------------------------------------------------------

_FORMATOS_FECHA = ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y", "%Y/%m/%d")


def _celda_a_texto(valor):
    """Texto de una celda tal como la escribió quien armó el archivo. Excel guarda los números
    como float: un CUIL tipeado sin formato llega como 20123456786.0 y no puede quedar con ese ".0"."""
    if valor is None:
        return ""
    if isinstance(valor, float):
        if valor != valor:  # NaN
            return ""
        return str(int(valor)) if valor.is_integer() else str(valor)
    if isinstance(valor, numbers.Integral) and not isinstance(valor, bool):
        return str(int(valor))
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%d/%m/%Y")
    return str(valor).strip()


def _parse_fecha(valor):
    """Devuelve un date, None si la celda está vacía, o levanta ValueError si no se entiende.

    Acepta lo que realmente trae un Excel: fecha real (datetime), número de serie de Excel y
    texto en DD/MM/AAAA (con / - o .) o AAAA-MM-DD, con o sin hora.
    """
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        resultado = valor.date()
    elif isinstance(valor, date):
        resultado = valor
    elif isinstance(valor, numbers.Real) and not isinstance(valor, bool):
        if not 1 <= valor <= 80000:
            raise ValueError(valor)
        resultado = (datetime(1899, 12, 30) + timedelta(days=float(valor))).date()
    else:
        texto = str(valor).strip().split(" ")[0].split("T")[0]
        resultado = None
        for formato in _FORMATOS_FECHA:
            try:
                resultado = datetime.strptime(texto, formato).date()
                break
            except ValueError:
                continue
        if resultado is None:
            raise ValueError(valor)
    if not 1900 <= resultado.year <= 2100:
        raise ValueError(valor)
    return resultado


def _parse_booleano(valor):
    return normalizar_texto(valor) in ("SI", "SÍ", "TRUE", "1", "X")


def _decodificar(contenido):
    """Los CSV que guarda Excel en español vienen en cp1252, no en UTF-8."""
    try:
        return contenido.decode("utf-8-sig")
    except UnicodeDecodeError:
        return contenido.decode("cp1252", errors="replace")


def _separador(texto):
    """Excel en español separa con ';'; otros programas con ','."""
    primera = texto.splitlines()[0] if texto else ""
    return max((";", ",", "\t"), key=primera.count)


def _leer_tabla(django_file):
    """Lee el archivo como tabla de texto/valores crudos. Levanta ArchivoInvalido si no se puede."""
    nombre = django_file.name.lower()
    try:
        if nombre.endswith(".csv"):
            texto = _decodificar(django_file.read())
            tabla = pd.read_csv(
                io.StringIO(texto), dtype=object, keep_default_na=False, sep=_separador(texto), engine="python",
            )
        else:
            tabla = pd.read_excel(django_file, dtype=object, keep_default_na=False)
    except Exception as error:
        raise ArchivoInvalido(
            "No pudimos leer el archivo. Verificá que sea un .xlsx o .csv válido y que no esté dañado "
            "(lo más seguro es descargar la plantilla y pegar ahí tus datos)."
        ) from error
    return tabla


def leer_archivo_empleados(django_file):
    """
    Lee un .xlsx/.csv subido por el cliente y devuelve un ResultadoLectura(filas, errores, columnas_ignoradas).

    filas: lista de dicts listos para instanciar/actualizar Empleado (con `numero_fila`).
    errores: lista de dicts {fila, mensaje} para las filas rechazadas.
    columnas_ignoradas: encabezados del archivo que no corresponden a ningún campo conocido.

    Levanta ArchivoInvalido cuando el archivo en sí no sirve (ilegible, sin las columnas
    obligatorias, vacío o demasiado grande): ahí no hay nada que cargar.
    """
    df = _leer_tabla(django_file)

    etiquetas = {campo: header.split(" (")[0] for header, campo, *_ in COLUMNAS}
    tipos_por_campo = {campo: tipo for _header, campo, _req, tipo in COLUMNAS}
    obligatorios = {campo for _header, campo, req, _tipo in COLUMNAS if req}

    # tolerante a variaciones menores de header (mayúsculas/espacios)
    columnas_normalizadas = {normalizar_texto(h): campo for h, campo, *_ in COLUMNAS}
    mapeo_reales, ignoradas = {}, []
    for col in df.columns:
        campo = columnas_normalizadas.get(normalizar_texto(col))
        if campo and campo not in mapeo_reales.values():
            mapeo_reales[col] = campo
        elif str(col).strip() and not str(col).startswith("Unnamed"):
            ignoradas.append(str(col).strip())
    df = df.rename(columns=mapeo_reales)

    faltantes = [etiquetas[c] for c in ("cuil", "nombre", "apellido") if c not in mapeo_reales.values()]
    if faltantes:
        raise ArchivoInvalido(
            "Al archivo le faltan columnas obligatorias: " + ", ".join(faltantes) + ". "
            "Revisá que la primera fila tenga los encabezados de la plantilla (CUIL, Nombre, Apellido, ...)."
        )

    # Las filas totalmente vacías (típicas al final de una planilla de Excel) no cuentan.
    df = df[[any(_celda_a_texto(v) for v in fila) for fila in df.itertuples(index=False)]]
    if df.empty:
        raise ArchivoInvalido("El archivo no tiene empleados para importar.")
    if len(df) > MAX_FILAS_IMPORTACION:
        raise ArchivoInvalido(
            f"El archivo tiene {len(df)} filas y el máximo por importación es {MAX_FILAS_IMPORTACION}. "
            "Dividilo en varios archivos."
        )

    from employees.models import Empleado  # acá y no arriba: este módulo se importa al cargar las apps

    largos = {
        f.name: f.max_length for f in Empleado._meta.get_fields()
        if getattr(f, "max_length", None) and f.name in tipos_por_campo and f.name != "cuil"
    }

    filas_validas, errores, cuils_vistos = [], [], set()

    for idx, row in df.iterrows():
        fila_num = idx + 2  # +1 por header, +1 porque idx es 0-based
        datos = {"numero_fila": fila_num}
        fila_errores = []

        for campo, tipo in tipos_por_campo.items():
            crudo = row.get(campo, "")
            if tipo in ("texto", "mail"):
                texto = _celda_a_texto(crudo)
                datos[campo] = normalizar_texto(texto) if campo in ("nombre", "apellido") else texto
                if tipo == "mail" and texto:
                    try:
                        validate_email(texto)
                    except ValidationError:
                        fila_errores.append(f"Mail inválido: '{texto}'")
                if campo in largos and len(datos[campo]) > largos[campo]:
                    fila_errores.append(f"'{etiquetas[campo]}' supera los {largos[campo]} caracteres")
            elif tipo == "fecha":
                try:
                    datos[campo] = _parse_fecha(crudo)
                except ValueError:
                    datos[campo] = None
                    fila_errores.append(
                        f"Fecha inválida en '{etiquetas[campo]}': {_celda_a_texto(crudo)!r} (usá DD/MM/AAAA)"
                    )
            elif tipo == "booleano":
                datos[campo] = _parse_booleano(crudo)
            elif tipo == "sexo":
                sexo = normalizar_texto(_celda_a_texto(crudo))
                datos[campo] = sexo[:1] if sexo in ("F", "M") else ""
            elif tipo == "cuil_ref":
                datos[campo] = normalizar_cuil(_celda_a_texto(crudo))
                if datos[campo] and not CUIL_RE.match(datos[campo]):
                    fila_errores.append(f"CUIL inválido en '{etiquetas[campo]}': debe tener 11 dígitos")

        datos["cuil"] = normalizar_cuil(_celda_a_texto(row.get("cuil", "")))

        for campo in obligatorios:
            if not datos.get(campo):
                fila_errores.append(f"Falta el campo obligatorio '{campo}'")

        if datos.get("cuil") and not CUIL_RE.match(datos["cuil"]):
            fila_errores.append("CUIL inválido: debe tener 11 dígitos")
        elif datos.get("cuil") in cuils_vistos:
            fila_errores.append(f"CUIL duplicado dentro del archivo: {datos['cuil']}")
        elif datos.get("cuil"):
            cuils_vistos.add(datos["cuil"])

        if fila_errores:
            errores.append({"fila": fila_num, "mensaje": "; ".join(fila_errores)})
        else:
            filas_validas.append(datos)

    return ResultadoLectura(filas_validas, errores, ignoradas)


def parse_archivo_empleados(django_file):
    """Versión corta de leer_archivo_empleados: devuelve solo (filas_validas, errores)."""
    resultado = leer_archivo_empleados(django_file)
    return resultado.filas, resultado.errores


def snapshot_empleados(implementacion):
    """Foto de los empleados para el historial de versiones de clients.
    Registrada en clients.services vía registrar_snapshot_builder (ver apps.py)."""
    return [
        {"CUIL": e.cuil, "Apellido y Nombre": f"{e.apellido}, {e.nombre}", "Legajo": e.legajo, "Puesto": e.puesto}
        for e in implementacion.empleados.all()
    ]
