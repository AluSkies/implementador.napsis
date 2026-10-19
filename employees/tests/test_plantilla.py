"""La plantilla de nómina que se descarga: más vistosa, con ayudas, listas desplegables e instrucciones."""

import io

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase

from employees.services import (
    COLUMNAS,
    COLOR_OBLIGATORIA,
    COLOR_OPCIONAL,
    generar_nomina_xlsx,
    generar_plantilla_xlsx,
    leer_archivo_empleados,
)


def abrir(buffer):
    return openpyxl.load_workbook(io.BytesIO(buffer.getvalue()))


class HojaDeEmpleadosTests(TestCase):
    def setUp(self):
        self.libro = abrir(generar_plantilla_xlsx())
        self.hoja = self.libro["Empleados"]

    def test_la_hoja_de_empleados_es_la_primera_y_la_que_queda_abierta(self):
        self.assertEqual(self.libro.sheetnames, ["Empleados", "Instrucciones"])
        self.assertEqual(self.libro.active.title, "Empleados")

    def test_los_encabezados_son_los_de_siempre(self):
        self.assertEqual([c.value for c in self.hoja[1]], [h for h, *_ in COLUMNAS])

    def test_las_obligatorias_van_en_violeta_y_las_opcionales_en_azul(self):
        colores = {c.value: c.fill.fgColor.rgb[-6:] for c in self.hoja[1]}
        for obligatoria in ("CUIL", "Nombre", "Apellido", "Fecha Ingreso (DD/MM/AAAA)"):
            self.assertEqual(colores[obligatoria], COLOR_OBLIGATORIA, obligatoria)
        for opcional in ("DNI", "Mail", "Puesto", "Reporta A (CUIL)"):
            self.assertEqual(colores[opcional], COLOR_OPCIONAL, opcional)

    def test_el_encabezado_va_en_blanco_y_negrita(self):
        celda = self.hoja["A1"]
        self.assertTrue(celda.font.bold)
        self.assertEqual(celda.font.color.rgb[-6:], "FFFFFF")

    def test_cada_encabezado_tiene_su_ayuda_con_un_ejemplo(self):
        for celda in self.hoja[1]:
            self.assertIsNotNone(celda.comment, celda.value)
            self.assertIn("Ejemplo:", celda.comment.text)
        self.assertTrue(self.hoja["A1"].comment.text.startswith("OBLIGATORIA."))
        self.assertTrue(self.hoja["B1"].comment.text.startswith("Opcional."))

    def test_la_fila_de_encabezados_queda_fija_y_con_filtros(self):
        self.assertEqual(self.hoja.freeze_panes, "A2")
        self.assertTrue(self.hoja.auto_filter.ref.startswith("A1:"))

    def test_sexo_y_los_si_no_tienen_lista_desplegable(self):
        listas = sorted(v.formula1 for v in self.hoja.data_validations.dataValidation)
        self.assertEqual(listas, ['"F,M"', '"SI,NO"', '"SI,NO"'])

    def test_la_plantilla_no_trae_filas_de_ejemplo_que_se_importarian_como_empleados(self):
        self.assertEqual(self.hoja.max_row, 1)  # primero esto: leer una celda la crea y cambiaría el resultado
        self.assertIsNone(self.hoja["A2"].value)


class HojaDeInstruccionesTests(TestCase):
    def setUp(self):
        self.hoja = abrir(generar_plantilla_xlsx())["Instrucciones"]
        self.texto = " ".join(str(c.value) for fila in self.hoja.iter_rows() for c in fila if c.value)

    def test_explica_como_completarla(self):
        for frase in ("una fila por persona", "No cambies los nombres de las columnas", "5.000 filas", "DD/MM/AAAA"):
            self.assertIn(frase, self.texto)

    def test_describe_cada_columna(self):
        for header, *_ in COLUMNAS:
            self.assertIn(header, self.texto)

    def test_marca_cuales_son_obligatorias(self):
        filas = {fila[0].value: fila[1].value for fila in self.hoja.iter_rows(min_row=10) if fila[0].value}
        self.assertEqual(filas["CUIL"], "Sí")
        self.assertEqual(filas["Mail"], "No")


class SigueSiendoImportableTests(TestCase):
    def test_una_fila_completada_en_la_plantilla_se_importa_sin_errores(self):
        libro = abrir(generar_plantilla_xlsx())
        hoja = libro["Empleados"]
        hoja.cell(2, 1, "20111111112")
        hoja.cell(2, 3, "Ana")
        hoja.cell(2, 4, "Gómez")
        hoja.cell(2, 6, "01/03/2018")
        salida = io.BytesIO()
        libro.save(salida)
        lectura = leer_archivo_empleados(SimpleUploadedFile("p.xlsx", salida.getvalue()), ingreso_obligatorio=True)
        self.assertEqual((len(lectura.filas), lectura.errores, lectura.columnas_ignoradas), (1, [], []))
        self.assertEqual(lectura.filas[0]["nombre"], "ANA")

    def test_la_nomina_que_descarga_el_equipo_tiene_los_mismos_encabezados(self):
        hoja = abrir(generar_nomina_xlsx([]))["Empleados"]
        self.assertEqual([c.value for c in hoja[1]], [h for h, *_ in COLUMNAS])
