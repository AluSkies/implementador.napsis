import io
from datetime import date, datetime
from unittest.mock import patch

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext

from clients.models import Cliente, Implementacion
from employees.forms import EmpleadoForm
from employees.models import Empleado, ImportBatch
from employees.services import ArchivoInvalido, generar_nomina_xlsx, leer_archivo_empleados
from onboarding import services

CABECERA = ["CUIL", "Nombre", "Apellido", "Fecha Ingreso (DD/MM/AAAA)", "Mail"]


def _xlsx(filas, cabecera=None, nombre="nomina.xlsx"):
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(cabecera or CABECERA)
    for fila in filas:
        hoja.append(fila)
    buffer = io.BytesIO()
    libro.save(buffer)
    return SimpleUploadedFile(nombre, buffer.getvalue())


def _csv(texto, codificacion="utf-8", nombre="nomina.csv"):
    return SimpleUploadedFile(nombre, texto.encode(codificacion), content_type="text/csv")


class LecturaDeFechasTests(TestCase):
    """Un Excel guarda las fechas como fechas reales, no como texto: antes se perdían en silencio."""

    def _fecha(self, valor):
        resultado = leer_archivo_empleados(_xlsx([[20111111112, "Juan", "Pérez", valor, ""]]))
        self.assertEqual(resultado.errores, [], resultado.errores)
        return resultado.filas[0]["fecha_ingreso"]

    def test_fecha_real_de_excel(self):
        self.assertEqual(self._fecha(datetime(2020, 3, 15)), date(2020, 3, 15))

    def test_fecha_como_texto_en_varios_formatos(self):
        for texto in ("15/03/2020", "15-03-2020", "2020-03-15", "5/3/2020"):
            with self.subTest(texto=texto):
                esperado = date(2020, 3, 5) if texto == "5/3/2020" else date(2020, 3, 15)
                self.assertEqual(self._fecha(texto), esperado)

    def test_numero_de_serie_de_excel(self):
        self.assertEqual(self._fecha(43905), date(2020, 3, 15))

    def test_celda_vacia_es_sin_fecha(self):
        self.assertIsNone(self._fecha(None))

    def test_una_fecha_que_no_se_entiende_es_un_error_de_fila_y_no_se_pierde_en_silencio(self):
        resultado = leer_archivo_empleados(_xlsx([[20111111112, "Juan", "Pérez", "mañana", ""]]))
        self.assertEqual(resultado.filas, [])
        self.assertIn("Fecha inválida en 'Fecha Ingreso'", resultado.errores[0]["mensaje"])
        self.assertIn("mañana", resultado.errores[0]["mensaje"])

    def test_un_anio_imposible_tambien_es_error(self):
        resultado = leer_archivo_empleados(_xlsx([[20111111112, "Juan", "Pérez", "15/03/0020", ""]]))
        self.assertEqual(len(resultado.errores), 1)


class LecturaDeArchivosRealesTests(TestCase):
    def test_cuil_numerico_de_excel_no_queda_con_decimales(self):
        for cuil in (20111111112, 20111111112.0, "20-11111111-2"):
            with self.subTest(cuil=cuil):
                resultado = leer_archivo_empleados(_xlsx([[cuil, "Juan", "Pérez", None, ""]]))
                self.assertEqual(resultado.filas[0]["cuil"], "20111111112")

    def test_csv_guardado_por_excel_en_espanol_cp1252_con_punto_y_coma(self):
        archivo = _csv("CUIL;Nombre;Apellido;Mail\n20111111112;José;Muñoz;jose@example.com\n", codificacion="cp1252")
        resultado = leer_archivo_empleados(archivo)
        self.assertEqual(resultado.errores, [])
        self.assertEqual((resultado.filas[0]["nombre"], resultado.filas[0]["apellido"]), ("JOSE", "MUNOZ"))

    def test_csv_utf8_con_punto_y_coma(self):
        resultado = leer_archivo_empleados(_csv("CUIL;Nombre;Apellido\n20111111112;Ana;Gómez\n"))
        self.assertEqual(resultado.errores, [])
        self.assertEqual(len(resultado.filas), 1)

    def test_un_archivo_roto_da_un_mensaje_claro_y_no_un_error_de_servidor(self):
        with self.assertRaisesMessage(ArchivoInvalido, "No pudimos leer el archivo"):
            leer_archivo_empleados(SimpleUploadedFile("nomina.xlsx", b"esto no es un excel"))

    def test_sin_las_columnas_obligatorias_lo_dice_una_vez_y_no_fila_por_fila(self):
        with self.assertRaisesMessage(ArchivoInvalido, "faltan columnas obligatorias: CUIL, Apellido"):
            leer_archivo_empleados(_xlsx([["Juan"]], cabecera=["Nombre"]))

    def test_las_filas_vacias_del_final_no_generan_errores(self):
        resultado = leer_archivo_empleados(_xlsx([[20111111112, "Juan", "Pérez", None, ""], [None] * 5, [None] * 5]))
        self.assertEqual(resultado.errores, [])
        self.assertEqual(len(resultado.filas), 1)

    def test_un_archivo_sin_empleados_se_rechaza(self):
        with self.assertRaisesMessage(ArchivoInvalido, "no tiene empleados"):
            leer_archivo_empleados(_xlsx([[None] * 5]))

    def test_las_columnas_desconocidas_se_informan(self):
        archivo = _xlsx([[20111111112, "Juan", "Pérez", None, "", "algo"]], cabecera=CABECERA + ["Talle de camisa"])
        self.assertEqual(leer_archivo_empleados(archivo).columnas_ignoradas, ["Talle de camisa"])

    def test_un_valor_mas_largo_que_el_campo_es_error_de_fila(self):
        cabecera = CABECERA + ["Telefono"]
        resultado = leer_archivo_empleados(_xlsx([[20111111112, "Juan", "Pérez", None, "", "9" * 60]], cabecera=cabecera))
        self.assertIn("'Telefono' supera los 50 caracteres", resultado.errores[0]["mensaje"])

    def test_un_mail_invalido_es_error_de_fila(self):
        resultado = leer_archivo_empleados(_xlsx([[20111111112, "Juan", "Pérez", None, "no-es-un-mail"]]))
        self.assertIn("Mail inválido", resultado.errores[0]["mensaje"])

    def test_el_cuil_de_la_jerarquia_tiene_que_tener_11_digitos(self):
        archivo = _xlsx([[20111111112, "Juan", "Pérez", None, "", "123"]], cabecera=CABECERA + ["Reporta A (CUIL)"])
        self.assertIn("CUIL inválido en 'Reporta A'", leer_archivo_empleados(archivo).errores[0]["mensaje"])

    def test_hay_un_tope_de_filas_por_archivo(self):
        archivo = _xlsx([[20111111110 + i, "Juan", "Pérez", None, ""] for i in range(4)])
        with patch("employees.services.MAX_FILAS_IMPORTACION", 3):
            with self.assertRaisesMessage(ArchivoInvalido, "el máximo por importación es 3"):
                leer_archivo_empleados(archivo)


class CargaDeLaImportacionTests(TestCase):
    def setUp(self):
        self.implementacion = Implementacion.objects.create(cliente=Cliente.objects.create(nombre="Empresa A"))

    def importar(self, archivo):
        return services.importar_archivo_empleados(self.implementacion, archivo, "ana")

    def test_carga_los_empleados_y_registra_la_importacion_sin_guardar_el_archivo(self):
        filas, errores, advertencias = self.importar(
            _xlsx([[20111111112, "Juan", "Pérez", datetime(2020, 3, 15), "j@x.com"], [20222222223, "Ana", "Gómez", None, ""]])
        )
        self.assertEqual((len(filas), errores, advertencias), (2, [], []))
        empleado = Empleado.objects.get(cuil="20111111112")
        self.assertEqual((empleado.apellido, empleado.fecha_ingreso), ("PEREZ", date(2020, 3, 15)))
        lote = ImportBatch.objects.get()
        self.assertEqual((lote.filas_importadas, lote.nombre_archivo), (2, "nomina.xlsx"))
        self.assertFalse(lote.archivo)  # los datos personales no quedan guardados en un archivo
        self.assertEqual(empleado.import_batch, lote)

    def test_volver_a_importar_actualiza_a_los_existentes_sin_duplicarlos(self):
        self.importar(_xlsx([[20111111112, "Juan", "Pérez", None, ""]]))
        self.importar(_xlsx([[20111111112, "Juan Carlos", "Pérez", None, "nuevo@x.com"]]))
        self.assertEqual(Empleado.objects.count(), 1)
        empleado = Empleado.objects.get()
        self.assertEqual((empleado.nombre, empleado.mail), ("JUAN CARLOS", "nuevo@x.com"))

    def test_la_jerarquia_se_resuelve_aunque_el_cuil_venga_con_guiones(self):
        cabecera = CABECERA + ["Reporta A (CUIL)", "Aprobador (CUIL)"]
        self.importar(_xlsx([
            [20111111112, "Juan", "Pérez", None, "", "20-22222222-3", "20-22222222-3"],
            [20222222223, "Ana", "Gómez", None, "", "", ""],
        ], cabecera=cabecera))
        juan = Empleado.objects.get(cuil="20111111112")
        self.assertEqual(juan.reporta_a.cuil, "20222222223")
        self.assertEqual(juan.aprobador.cuil, "20222222223")

    def test_un_cuil_de_jerarquia_que_no_existe_se_avisa_y_no_pisa_la_relacion(self):
        cabecera = CABECERA + ["Reporta A (CUIL)"]
        _filas, _errores, advertencias = self.importar(_xlsx([
            [20111111112, "Juan", "Pérez", None, "", "20999999990"],
        ], cabecera=cabecera))
        self.assertIn("Fila 2", advertencias[0])
        self.assertIn("20999999990", advertencias[0])
        self.assertIsNone(Empleado.objects.get().reporta_a)

    def test_las_columnas_ignoradas_llegan_como_advertencia(self):
        _f, _e, advertencias = self.importar(
            _xlsx([[20111111112, "Juan", "Pérez", None, "", "x"]], cabecera=CABECERA + ["Talle de camisa"])
        )
        self.assertIn("Talle de camisa", advertencias[0])

    def test_las_filas_con_errores_no_impiden_cargar_las_demas(self):
        filas, errores, _a = self.importar(_xlsx([
            [20111111112, "Juan", "Pérez", None, ""], ["123", "Sin", "Cuil", None, ""],
        ]))
        self.assertEqual((len(filas), len(errores)), (1, 1))
        self.assertEqual(Empleado.objects.count(), 1)

    def test_es_todo_o_nada_si_algo_falla_a_mitad_de_la_carga(self):
        with patch("onboarding.services._resolver_jerarquia_importada", side_effect=RuntimeError("falló la base")):
            with self.assertRaises(RuntimeError):
                self.importar(_xlsx([[20111111112, "Juan", "Pérez", None, ""]]))
        self.assertEqual((Empleado.objects.count(), ImportBatch.objects.count()), (0, 0))

    def test_un_archivo_invalido_no_toca_nada(self):
        with self.assertRaises(ArchivoInvalido):
            self.importar(SimpleUploadedFile("nomina.xlsx", b"roto"))
        self.assertEqual((Empleado.objects.count(), ImportBatch.objects.count()), (0, 0))
        self.assertFalse(self.implementacion.versiones.exists())

    def test_una_nomina_grande_se_carga_con_pocas_consultas(self):
        """Antes eran ~9 consultas por fila: con miles de empleados, el pedido web se cortaba."""
        archivo = _xlsx([[20000000000 + i, "Juan", f"Pérez{i}", None, ""] for i in range(600)])
        with CaptureQueriesContext(connection) as consultas:
            filas, errores, _a = self.importar(archivo)
        self.assertEqual((len(filas), errores), (600, []))
        self.assertLess(len(consultas), 40)


class NominaDescargadaTests(TestCase):
    def test_un_texto_que_parece_formula_sale_como_texto(self):
        implementacion = Implementacion.objects.create(cliente=Cliente.objects.create(nombre="Empresa A"))
        Empleado.objects.create(
            implementacion=implementacion, cuil="20111111112", nombre="JUAN", apellido="PEREZ",
            domicilio='=HYPERLINK("http://malo.example","clic")',
        )
        libro = openpyxl.load_workbook(generar_nomina_xlsx(implementacion.empleados.all()))
        celdas = [c for fila in libro.active.iter_rows(min_row=2) for c in fila if c.value]
        formula = next(c for c in celdas if str(c.value).startswith("=HYPERLINK"))
        self.assertEqual(formula.data_type, "s")  # texto, no fórmula


class EmpleadoFormTests(TestCase):
    def setUp(self):
        self.implementacion = Implementacion.objects.create(cliente=Cliente.objects.create(nombre="Empresa A"))

    def form(self, **datos):
        base = {"cuil": "20-11111111-2", "nombre": "juan", "apellido": "pérez"}
        return EmpleadoForm({**base, **datos}, implementacion=self.implementacion)

    def test_normaliza_cuil_y_nombres_como_la_importacion(self):
        formulario = self.form()
        self.assertTrue(formulario.is_valid(), formulario.errors)
        self.assertEqual(formulario.cleaned_data["cuil"], "20111111112")
        self.assertEqual((formulario.cleaned_data["nombre"], formulario.cleaned_data["apellido"]), ("JUAN", "PEREZ"))

    def test_rechaza_un_cuil_que_no_tiene_11_digitos(self):
        for cuil in ("abc", "123", "20-1111-2"):
            with self.subTest(cuil=cuil):
                self.assertIn("11 dígitos", " ".join(self.form(cuil=cuil).errors["cuil"]))

    def test_rechaza_un_cuil_repetido_aunque_venga_con_guiones(self):
        Empleado.objects.create(implementacion=self.implementacion, cuil="20111111112", nombre="A", apellido="B")
        self.assertIn("Ya hay un empleado", " ".join(self.form().errors["cuil"]))

    def test_al_editar_no_se_choca_con_si_mismo(self):
        existente = Empleado.objects.create(implementacion=self.implementacion, cuil="20111111112", nombre="A", apellido="B")
        formulario = EmpleadoForm(
            {"cuil": "20111111112", "nombre": "A", "apellido": "C"}, instance=existente, implementacion=self.implementacion,
        )
        self.assertTrue(formulario.is_valid(), formulario.errors)

    def test_el_doble_clic_no_da_error_de_servidor(self):
        """Dos pedidos validan antes de que cualquiera guarde: el segundo choca recién en la base
        y tiene que volver como error de campo, no como página de error."""
        primero, segundo = self.form(), self.form()
        self.assertTrue(primero.is_valid() and segundo.is_valid())
        self.assertTrue(services.agregar_empleado(self.implementacion, primero, "ana"))
        self.assertFalse(services.agregar_empleado(self.implementacion, segundo, "ana"))
        self.assertIn("Ya hay un empleado", " ".join(segundo.errors["cuil"]))
        self.assertEqual(Empleado.objects.count(), 1)


class ReimportarLaMismaPlanillaTests(TestCase):
    def test_no_reescribe_a_los_empleados_que_no_cambiaron(self):
        implementacion = Implementacion.objects.create(cliente=Cliente.objects.create(nombre="Empresa A"))
        filas = [[20000000000 + i, "Juan", f"Pérez{i}", None, ""] for i in range(20)]
        services.importar_archivo_empleados(implementacion, _xlsx(filas), "ana")
        antes = dict(Empleado.objects.values_list("cuil", "actualizado"))

        filas[0][1] = "Juan Carlos"  # uno solo cambia
        services.importar_archivo_empleados(implementacion, _xlsx(filas), "ana")

        despues = dict(Empleado.objects.values_list("cuil", "actualizado"))
        cambiados = [cuil for cuil in antes if antes[cuil] != despues[cuil]]
        self.assertEqual(cambiados, ["20000000000"])
        self.assertEqual(Empleado.objects.get(cuil="20000000000").nombre, "JUAN CARLOS")
