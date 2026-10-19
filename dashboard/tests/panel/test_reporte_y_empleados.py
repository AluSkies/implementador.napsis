import io
from unittest.mock import patch

import openpyxl
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from clients.models import Cliente, Implementacion, ImplementacionModulo, Modulo
from employees.models import Empleado

from dashboard.tests.base import BasePanel


class ReporteExportableTests(BasePanel):
    def setUp(self):
        super().setUp()
        cliente = Cliente.objects.create(nombre="Empresa A", razon_social="Empresa A S.A.", cuits="30500010912")
        self.implementacion = Implementacion.objects.create(
            cliente=cliente, creado_por=self.ana, implementador_asignado=self.ana,
            mail_contacto="contacto@a.com", telefono="011-1", paso_actual=4,
        )
        Empleado.objects.create(implementacion=self.implementacion, cuil="20111111112", nombre="J", apellido="P")
        modulo = Modulo.objects.create(nombre="Fichada")
        ImplementacionModulo.objects.create(implementacion=self.implementacion, modulo=modulo, activo=True)

    def descargar(self):
        respuesta = self.client.get(reverse("dashboard:exportar"))
        self.assertEqual(respuesta.status_code, 200)
        libro = openpyxl.load_workbook(io.BytesIO(b"".join(respuesta.streaming_content)))
        return respuesta, [[c.value for c in fila] for fila in libro.active.iter_rows()]

    def test_el_panel_tiene_el_boton_al_pie(self):
        respuesta = self.client.get(reverse("dashboard:home"))
        self.assertContains(respuesta, reverse("dashboard:exportar"))
        self.assertContains(respuesta, "Exportar (.xlsx)")

    def test_descarga_un_xlsx_con_una_fila_por_implementacion(self):
        respuesta, filas = self.descargar()
        self.assertIn("implementaciones_", respuesta["Content-Disposition"])
        self.assertTrue(respuesta["Content-Disposition"].endswith('.xlsx"'))
        cabecera = filas[0]
        for columna in ("Empresa", "Razón social", "CUIT(s)", "Implementador asignado", "Fecha de solicitud",
                        "Fecha de implementación", "Mail de contacto"):
            self.assertIn(columna, cabecera)
        self.assertEqual(len(filas), 2)

    def test_la_fila_trae_los_datos_de_la_implementacion(self):
        _r, filas = self.descargar()
        fila = dict(zip(filas[0], filas[1], strict=True))
        self.assertEqual(fila["Empresa"], "Empresa A")
        self.assertEqual(fila["Razón social"], "Empresa A S.A.")
        self.assertEqual(fila["CUIT(s)"], "30500010912")
        self.assertEqual(fila["Implementador asignado"], "Ana Paz")
        self.assertEqual(fila["Creada por"], "ana")
        self.assertEqual(fila["Mail de contacto"], "contacto@a.com")
        self.assertEqual(fila["Módulos"], "Fichada")
        self.assertEqual((fila["Empleados"], fila["Sucursales"]), (1, 0))
        self.assertEqual(fila["Paso actual"], 4)
        self.assertEqual(fila["Estado"], "Pendiente - A completar por cliente")

    def test_la_fecha_de_implementacion_es_real_si_esta_cumplida_y_estimada_si_no(self):
        from datetime import date

        self.implementacion.fecha_salida_productivo = date(2026, 12, 1)
        self.implementacion.save()
        _r, filas = self.descargar()
        fila = dict(zip(filas[0], filas[1], strict=True))
        self.assertEqual((fila["Fecha de implementación"], fila["Tipo de fecha de implementación"]), ("01/12/2026", "Estimada"))

        self.implementacion.marcar_cumplida()
        _r, filas = self.descargar()
        fila = dict(zip(filas[0], filas[1], strict=True))
        self.assertEqual(fila["Tipo de fecha de implementación"], "Real")
        self.assertEqual(fila["Estado"], "Completo")

    def test_incluye_todas_sin_importar_el_estado(self):
        otra = Implementacion.objects.create(cliente=Cliente.objects.create(nombre="Empresa B"))
        otra.anular()
        _r, filas = self.descargar()
        self.assertEqual(len(filas), 3)

    def test_un_dato_que_parece_formula_sale_como_texto(self):
        self.implementacion.cliente.razon_social = "=1+1"
        self.implementacion.cliente.save()
        respuesta = self.client.get(reverse("dashboard:exportar"))
        libro = openpyxl.load_workbook(io.BytesIO(b"".join(respuesta.streaming_content)))
        celda = next(c for c in libro.active[2] if c.value == "=1+1")
        self.assertEqual(celda.data_type, "s")

    def test_exige_ser_del_equipo(self):
        self.client.force_login(self.cliente_user)
        self.assertEqual(self.client.get(reverse("dashboard:exportar")).status_code, 302)
        self.client.logout()
        self.assertEqual(self.client.get(reverse("dashboard:exportar")).status_code, 302)


class EmpleadosEnElPanelTests(BasePanel):
    def setUp(self):
        super().setUp()
        self.implementacion = Implementacion.objects.create(cliente=Cliente.objects.create(nombre="Empresa A"))
        self.empleado = Empleado.objects.create(
            implementacion=self.implementacion, cuil="20111111112", nombre="JUAN", apellido="PEREZ",
        )
        self.url = reverse("dashboard:implementacion_detalle", kwargs={"pk": self.implementacion.pk})

    def test_cada_empleado_tiene_su_boton_de_editar(self):
        self.assertContains(self.client.get(self.url), f"?editar_empleado={self.empleado.pk}#empleados")

    def test_editar_abre_el_formulario_con_sus_datos(self):
        respuesta = self.client.get(self.url + f"?editar_empleado={self.empleado.pk}")
        self.assertContains(respuesta, "Editando a PEREZ, JUAN")
        self.assertEqual(respuesta.context["empleado_edit_form"].instance, self.empleado)

    def test_guarda_los_cambios_y_deja_historial(self):
        respuesta = self.client.post(self.url, {
            "accion": "editar_empleado", "empleado_id": self.empleado.pk,
            "cuil": "20-11111111-2", "nombre": "Juan", "apellido": "Pérez", "puesto": "Analista",
        })
        self.assertEqual(respuesta.status_code, 302)
        self.empleado.refresh_from_db()
        self.assertEqual((self.empleado.cuil, self.empleado.puesto), ("20111111112", "Analista"))
        self.assertEqual(self.implementacion.versiones.get(seccion="empleados").editado_por, "ana")

    def test_un_cuil_repetido_vuelve_con_el_error(self):
        Empleado.objects.create(implementacion=self.implementacion, cuil="20222222223", nombre="ANA", apellido="GOMEZ")
        respuesta = self.client.post(self.url, {
            "accion": "editar_empleado", "empleado_id": self.empleado.pk,
            "cuil": "20222222223", "nombre": "Juan", "apellido": "Pérez",
        })
        self.assertEqual(respuesta.status_code, 200)
        self.assertContains(respuesta, "Ya hay un empleado cargado con ese CUIL")

    def test_un_archivo_roto_se_muestra_como_error_del_formulario(self):
        respuesta = self.client.post(self.url, {
            "accion": "importar_empleados", "archivo": SimpleUploadedFile("nomina.xlsx", b"no es excel"),
        })
        self.assertEqual(respuesta.status_code, 200)
        self.assertContains(respuesta, "No pudimos leer el archivo")

    def test_importar_informa_el_resultado(self):
        libro = openpyxl.Workbook()
        libro.active.append(["CUIL", "Nombre", "Apellido"])
        libro.active.append([20222222223, "Ana", "Gómez"])
        libro.active.append(["123", "Mal", "Cuil"])
        buffer = io.BytesIO()
        libro.save(buffer)
        respuesta = self.client.post(self.url, {
            "accion": "importar_empleados", "archivo": SimpleUploadedFile("nomina.xlsx", buffer.getvalue()),
        }, follow=True)
        textos = " ".join(str(m) for m in respuesta.context["messages"])
        self.assertIn("Se importaron 1 empleados. 1 filas tuvieron errores", textos)
        self.assertIn("fila 3", textos)
        self.assertTrue(Empleado.objects.filter(cuil="20222222223").exists())


class AltaAtomicaTests(BasePanel):
    """Si algo falla a mitad del alta no tiene que quedar una empresa o un usuario suelto."""

    def test_si_falla_al_guardar_la_implementacion_no_queda_nada(self):
        with patch("onboarding.acciones.services.guardar_modulos", side_effect=RuntimeError("base caída")):
            with self.assertLogs("django.request", "ERROR"), self.assertRaises(RuntimeError):
                self.crear(nombre="Empresa Nueva", username="nueva", mail_cliente="gerencia@nueva.com")
        self.assertFalse(Cliente.objects.exists())
        self.assertFalse(Implementacion.objects.exists())
        self.assertFalse(User.objects.filter(username="nueva").exists())

    def test_si_falla_el_alta_no_se_manda_la_invitacion(self):
        from django.core import mail

        with patch("onboarding.acciones.services.guardar_modulos", side_effect=RuntimeError("base caída")):
            with self.assertLogs("django.request", "ERROR"), self.assertRaises(RuntimeError):
                self.crear(nombre="Empresa Nueva", username="nueva", mail_cliente="gerencia@nueva.com")
        self.assertEqual(mail.outbox, [])

    def test_un_alta_exitosa_sigue_creando_todo(self):
        self.crear(nombre="Empresa Nueva", username="nueva")
        self.assertEqual((Cliente.objects.count(), Implementacion.objects.count()), (1, 1))
        self.assertTrue(User.objects.filter(username="nueva").exists())


class BotonesAntiDobleClicTests(BasePanel):
    def test_la_pagina_deshabilita_los_botones_al_enviar(self):
        pagina = self.client.get(reverse("dashboard:home")).content.decode()
        self.assertIn("document.addEventListener('submit'", pagina)
        self.assertIn("b.disabled = true", pagina)
