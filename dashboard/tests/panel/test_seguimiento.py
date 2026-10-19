"""El seguimiento en el panel: marca de «sin movimiento» y comentarios del equipo."""

from datetime import timedelta

import openpyxl
from django.urls import reverse
from django.utils import timezone

from clients.models import Cliente, Implementacion
from clients.tests.modelos.test_sin_movimiento import evento_de_hace, implementacion_de_hace
from dashboard.reportes import generar_reporte_implementaciones_xlsx
from dashboard.tests.base import BasePanel


def tabla_del_listado(pagina):
    """Solo la tabla de implementaciones (los nombres también están en el desplegable del alta)."""
    return pagina.split("<tbody>")[1].split("</tbody>")[0]


def fila_de(pagina, nombre):
    """El HTML de la fila del listado que corresponde a una empresa."""
    tabla = tabla_del_listado(pagina)
    return tabla[tabla.index(nombre):].split("</tr>")[0]


class SinMovimientoEnElPanelTests(BasePanel):
    def setUp(self):
        super().setUp()
        self.quieta = implementacion_de_hace(10, "Quieta S.A.")
        self.activa = implementacion_de_hace(10, "Activa S.A.")
        evento_de_hace(self.activa, 1)
        self.home = reverse("dashboard:home")

    def test_el_listado_marca_solo_la_que_esta_quieta(self):
        pagina = self.client.get(self.home).content.decode()
        self.assertIn("Implementación sin movimiento", fila_de(pagina, "Quieta S.A."))
        self.assertNotIn("Implementación sin movimiento", fila_de(pagina, "Activa S.A."))

    def test_la_marca_no_tapa_el_estado_real(self):
        pagina = self.client.get(self.home).content.decode()
        self.assertIn("Pendiente - A completar por cliente", fila_de(pagina, "Quieta S.A."))

    def test_se_puede_filtrar(self):
        pagina = self.client.get(self.home, {"estado": "sin_movimiento"}).content.decode()
        self.assertIn("Quieta S.A.", tabla_del_listado(pagina))
        self.assertNotIn("Activa S.A.", tabla_del_listado(pagina))
        self.assertIn("Implementación sin movimiento</option>", pagina)

    def test_la_ficha_avisa_hace_cuanto_y_que_la_saca(self):
        pagina = self.client.get(reverse("dashboard:implementacion_detalle", kwargs={"pk": self.quieta.pk}))
        self.assertContains(pagina, "Implementación sin movimiento:")
        self.assertContains(pagina, "hace 10 días")
        self.assertContains(pagina, "Cualquier acción, comentario o envío de documentación")

    def test_la_ficha_de_una_activa_no_avisa(self):
        pagina = self.client.get(reverse("dashboard:implementacion_detalle", kwargs={"pk": self.activa.pk}))
        self.assertNotContains(pagina, "Implementación sin movimiento:")

    def test_comentar_desde_la_ficha_la_saca_de_ahi(self):
        url = reverse("dashboard:implementacion_detalle", kwargs={"pk": self.quieta.pk})
        self.client.post(url, {"accion": "agregar_comentario", "texto": "Hablé con el contador."})
        self.assertNotContains(self.client.get(url), "Implementación sin movimiento:")
        pagina = self.client.get(self.home, {"estado": "sin_movimiento"}).content.decode()
        self.assertNotIn("Quieta S.A.", tabla_del_listado(pagina))

    def test_una_cumplida_no_se_marca(self):
        implementacion_de_hace(30, "Cumplida S.A.", estado=Implementacion.ESTADO_CUMPLIDA)
        pagina = self.client.get(self.home, {"estado": "sin_movimiento"}).content.decode()
        self.assertNotIn("Cumplida S.A.", tabla_del_listado(pagina))

    def test_el_reporte_tiene_la_ultima_actividad_y_la_marca(self):
        hoja = openpyxl.load_workbook(generar_reporte_implementaciones_xlsx()).active
        encabezados = [hoja.cell(1, c).value for c in range(1, hoja.max_column + 1)]
        filas = {
            hoja.cell(r, 1).value: dict(zip(encabezados, [hoja.cell(r, c).value for c in range(1, hoja.max_column + 1)], strict=True))
            for r in range(2, hoja.max_row + 1)
        }
        self.assertEqual(filas["Quieta S.A."]["Sin movimiento"], "Sí")
        self.assertEqual(filas["Activa S.A."]["Sin movimiento"], "No")
        ayer = timezone.localtime(timezone.now() - timedelta(days=1)).strftime("%d/%m/%Y")
        self.assertEqual(filas["Activa S.A."]["Última actividad"], ayer)


class ComentariosEnLaFichaTests(BasePanel):
    def setUp(self):
        super().setUp()
        self.implementacion = Implementacion.objects.create(cliente=Cliente.objects.create(nombre="Empresa A"))
        self.url = reverse("dashboard:implementacion_detalle", kwargs={"pk": self.implementacion.pk})

    def comentar(self, texto):
        return self.client.post(self.url, {"accion": "agregar_comentario", "texto": texto})

    def test_la_ficha_ofrece_comentar_y_explica_que_es_privado(self):
        pagina = self.client.get(self.url)
        self.assertContains(pagina, "Comentarios del equipo")
        self.assertContains(pagina, "el cliente nunca las ve")
        self.assertContains(pagina, "Todavía no hay comentarios.")

    def test_se_muestra_con_autor_y_el_mas_nuevo_primero(self):
        self.comentar("Primera nota")
        self.comentar("Segunda nota")
        pagina = self.client.get(self.url).content.decode()
        self.assertLess(pagina.index("Segunda nota"), pagina.index("Primera nota"))
        self.assertIn("Ana Paz", pagina)

    def test_un_comentario_vacio_avisa_y_no_guarda(self):
        respuesta = self.client.post(self.url, {"accion": "agregar_comentario", "texto": "  "}, follow=True)
        self.assertContains(respuesta, "Escribí el comentario.")
        self.assertFalse(self.implementacion.comentarios.exists())

    def test_el_html_del_comentario_se_escapa(self):
        self.comentar("<script>alert(1)</script>")
        self.assertNotContains(self.client.get(self.url), "<script>alert(1)</script>")

    def test_los_saltos_de_linea_se_respetan(self):
        self.comentar("línea uno\nlínea dos")
        self.assertContains(self.client.get(self.url), "línea uno<br>línea dos")
