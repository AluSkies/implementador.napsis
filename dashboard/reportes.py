"""Reporte exportable con todas las implementaciones (botón "Exportar" al pie del panel)."""

import io

import openpyxl
from django.db.models import Count
from django.utils import timezone

from clients.models import Implementacion
from clients.xlsx import escribir_celda

COLUMNAS = [
    ("Empresa", 28), ("Razón social", 32), ("CUIT(s)", 26), ("Estado", 32), ("Paso actual", 12),
    ("Implementador asignado", 24), ("Creada por", 18), ("Mail de contacto", 30), ("Teléfono", 18),
    ("Fecha de solicitud", 18), ("Fecha de envío del cliente", 22), ("Fecha de implementación", 22),
    ("Tipo de fecha de implementación", 24), ("Módulos", 40), ("Empleados", 12), ("Sucursales", 12),
    ("Pidió contacto sin atender", 18), ("Fecha de validación", 20), ("Rondas de devolución", 14),
]


def _fecha(valor):
    return timezone.localtime(valor).strftime("%d/%m/%Y") if valor else ""


def generar_reporte_implementaciones_xlsx():
    """Un .xlsx con una fila por implementación (todas, sin filtros), de la más nueva a la más
    vieja. La fecha de implementación es la real si ya está cumplida, o la estimada si no."""
    implementaciones = (
        Implementacion.objects.select_related("cliente", "implementador_asignado", "creado_por")
        .prefetch_related("modulos__modulo")
        .annotate(
            cant_empleados=Count("empleados", distinct=True), cant_sucursales=Count("sucursales", distinct=True),
            cant_rondas=Count("rondas", distinct=True),
        )
        .order_by("-creado")
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Implementaciones"
    for col, (titulo, ancho) in enumerate(COLUMNAS, start=1):
        celda = ws.cell(row=1, column=col, value=titulo)
        celda.font = openpyxl.styles.Font(bold=True)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = ancho
    ws.freeze_panes = "A2"

    for fila, imp in enumerate(implementaciones, start=2):
        if imp.estado == Implementacion.ESTADO_CUMPLIDA:
            fecha_impl, tipo_fecha = _fecha(imp.cumplida_en), "Real"
        elif imp.fecha_salida_productivo:
            fecha_impl, tipo_fecha = imp.fecha_salida_productivo.strftime("%d/%m/%Y"), "Estimada"
        else:
            fecha_impl, tipo_fecha = "", ""
        valores = [
            imp.cliente.nombre, imp.cliente.razon_social, imp.cliente.cuits,
            imp.marca_de_devolucion or imp.estado_visual_label,
            imp.paso_actual if imp.estado_visual == Implementacion.VISUAL_A_COMPLETAR else "",
            imp.asignado_nombre, imp.creado_por.get_username() if imp.creado_por else "",
            imp.mail_contacto, imp.telefono, _fecha(imp.creado), _fecha(imp.enviado), fecha_impl, tipo_fecha,
            ", ".join(m.modulo.nombre for m in imp.modulos.all() if m.activo),
            imp.cant_empleados, imp.cant_sucursales,
            "Sí" if imp.revision_solicitada and not imp.revision_atendida_en else "No",
            _fecha(imp.validada_en), imp.cant_rondas,
        ]
        for col, valor in enumerate(valores, start=1):
            escribir_celda(ws, fila, col, valor)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
